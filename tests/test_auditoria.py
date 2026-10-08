"""
test_auditoria.py
=============================================================================
Suite de pruebas requerida para validar la validación compuesta (Correo + Nombre),
la normalización fonética/ortográfica y los criterios de riesgo tripartita:
1. Riesgos Activos (CRÍTICO - RIESGO ACTIVO, ALTO - CUENTA HUÉRFANA) -> Pestaña 2
2. Incidentes Resueltos / Pasados (MEDIO - INCIDENTE PASADO) -> Pestaña 3
3. Reingresos y Casos OK (POSIBLE REINGRESO, CONFORME) -> Pestaña 4 y Pestaña 1
=============================================================================
"""

import sys
from pathlib import Path
import openpyxl
import pandas as pd
import pytest

# Permitir importación del módulo principal al ejecutar directamente desde tests/
ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from auditor_bajas import (
    clean_identifier_series,
    clean_name_series,
    consolidate_logins,
    create_glosario_sheet,
    deduplicate_universo,
    execute_audit,
    export_to_excel,
    normalize_text_name,
    run_audit_pipeline,
    select_closest_login_event,
)


def test_caso_1_mismo_correo_nombres_distintos():
    """
    Caso 1 (Mismo correo, nombres distintos):
    Dos usuarios comparten o reúsan el correo soporte@empresa.com:
    - En Universo: CARLOS LOPEZ fue dado de baja el 10/08/2026 con correo soporte@empresa.com.
    - En Reporte: El correo soporte@empresa.com está asignado a JUAN PEREZ con logins en septiembre 2026.
    Resultado esperado:
    - No vincular arbitrariamente el login de Juan Perez a Carlos Lopez.
    - Carlos Lopez no debe tener login atribuido (ULTIMO_LOGIN_DETECTADO = 'Sin registro').
    - Se debe marcar DISCREPANCIA_IDENTIDAD = True para Carlos Lopez.
    - Carlos Lopez debe resultar 'OK' y 'CONFORME' (no es acceso post-baja de Carlos).
    """
    df_uni = pd.DataFrame({
        "CLA_TRAB": [101],
        "CLA_LOGIN": ["CLOPEZ"],
        "NOMBRE": ["CARLOS LOPEZ"],
        "FECHA_BAJA": ["10/08/2026"],
        "CORREO": ["soporte@empresa.com"],
    })

    df_rep = pd.DataFrame({
        "Full Name": ["JUAN PEREZ"],
        "Username": ["soporte@empresa.com"],
        "Active": ["1"],
        "Last Login": ["29/09/2026 10:00 AM"],
    })

    col_map = {
        "name": "Full Name",
        "user": "Username",
        "active": "Active",
        "login": "Last Login",
        "baja_rep": None,
    }

    consolidated = consolidate_logins(df_rep, col_map)
    df_completo, df_hallazgos, metrics = execute_audit(
        df_universo=df_uni,
        baja_col="FECHA_BAJA",
        email_col="CORREO",
        nombre_col="NOMBRE",
        consolidated_logins=consolidated,
    )

    carlos = df_completo.iloc[0]
    # No debe haberse cruzado con Juan Perez
    assert carlos["ULTIMO_LOGIN_DETECTADO"] == "Sin registro"
    # Debe haberse detectado la discrepancia de identidad
    assert bool(carlos["DISCREPANCIA_IDENTIDAD"]) is True
    # Al no haber acceso post-baja de Carlos ni cuenta activa de Carlos, es OK / CONFORME
    assert carlos["ESTATUS_AUDITORIA"] == "OK"
    assert carlos["CATEGORIA_RIESGO"] == "CONFORME"
    assert len(df_hallazgos) == 0


def test_caso_2_acentos_y_espacios_match_perfecto():
    """
    Caso 2 (Acentos y espacios con cuenta desactivada Active == 0):
    JOSE RICARDO FRÍAS MIRELES en Universo vs Jose Ricardo Frias Mireles en Reporte
    deben hacer match perfecto a través de la normalización.
    Al tener Active == 0 y login posterior a la baja, se clasifica como:
    INCIDENTE RESUELTO / MEDIO - INCIDENTE PASADO.
    """
    n_uni = "JOSE RICARDO FRÍAS MIRELES"
    n_rep = "  Jose   Ricardo   Frias   Mireles  "

    # Prueba directa de la función de normalización
    assert normalize_text_name(n_uni) == normalize_text_name(n_rep)
    assert normalize_text_name(n_uni) == "jose ricardo frias mireles"

    # Prueba en pipeline de cruce
    df_uni = pd.DataFrame({
        "CLA_TRAB": [501],
        "CLA_LOGIN": ["JFRIAS"],
        "NOMBRE": [n_uni],
        "FECHA_BAJA": ["03/09/2026"],
        "CORREO": ["jfrias@empresa.com"],
    })

    df_rep = pd.DataFrame({
        "Full Name": [n_rep],
        "Username": ["jfrias@empresa.com"],
        "Active": ["0"],
        "Last Login": ["15/09/2026 09:00"],
    })

    col_map = {
        "name": "Full Name",
        "user": "Username",
        "active": "Active",
        "login": "Last Login",
        "baja_rep": None,
    }

    consolidated = consolidate_logins(df_rep, col_map)
    df_completo, df_hallazgos, metrics = execute_audit(
        df_universo=df_uni,
        baja_col="FECHA_BAJA",
        email_col="CORREO",
        nombre_col="NOMBRE",
        consolidated_logins=consolidated,
    )

    frias = df_completo.iloc[0]
    # Debe cruzar perfectamente
    assert frias["ULTIMO_LOGIN_DETECTADO"] == "15/09/2026"
    assert bool(frias["DISCREPANCIA_IDENTIDAD"]) is False
    assert frias["ESTATUS_AUDITORIA"] == "INCIDENTE RESUELTO"
    assert frias["CATEGORIA_RIESGO"] == "MEDIO - INCIDENTE PASADO"
    assert frias["DIAS_POST_BAJA"] == 12
    # Al estar ya desactivada la cuenta (Active == 0), no es riesgo activo
    assert len(df_hallazgos) == 0


def test_caso_3_cuenta_activa_post_baja():
    """
    Caso 3 (Cuenta activa post-baja):
    Active == 1 con baja confirmada pero sin login posterior a la baja
    debe disparar REVISAR / ALTO - CUENTA HUÉRFANA.
    """
    df_uni = pd.DataFrame({
        "CLA_TRAB": [201],
        "CLA_LOGIN": ["APAREDE"],
        "NOMBRE": ["ANETTE MICHELLE PAREDES RUIZ"],
        "FECHA_BAJA": ["27/05/2026"],
        "CORREO": ["aparedes@empresa.com"],
    })

    df_rep = pd.DataFrame({
        "Full Name": ["Anette Michelle Paredes Ruiz"],
        "Username": ["aparedes@empresa.com"],
        "Active": ["1"],
        # Login anterior a la fecha de baja
        "Last Login": ["26/05/2026 14:00"],
    })

    col_map = {
        "name": "Full Name",
        "user": "Username",
        "active": "Active",
        "login": "Last Login",
        "baja_rep": None,
    }

    consolidated = consolidate_logins(df_rep, col_map)
    df_completo, df_hallazgos, metrics = execute_audit(
        df_universo=df_uni,
        baja_col="FECHA_BAJA",
        email_col="CORREO",
        nombre_col="NOMBRE",
        consolidated_logins=consolidated,
    )

    paredes = df_completo.iloc[0]
    assert paredes["ESTATUS_CUENTA_REPORTE"] == "Activa"
    assert paredes["ESTATUS_AUDITORIA"] == "REVISAR"
    assert paredes["CATEGORIA_RIESGO"] == "ALTO - CUENTA HUÉRFANA"
    assert pd.isna(paredes["DIAS_POST_BAJA"])
    assert len(df_hallazgos) == 1


def test_criterios_validacion_matriz_riesgo(tmp_path):
    """
    [V] - VALIDATION CRITERIA:
    Verifica los criterios obligatorios de la misión:
    1. Usuario con login el 15/09/2026, baja el 01/09/2026 y Active == 0
       se clasifica estrictamente como INCIDENTE RESUELTO / MEDIO - INCIDENTE PASADO
       y se archiva en la Pestaña 3 (Incidentes_Pasados).
    2. Usuario con login el 15/09/2026, baja el 01/09/2026 y Active == 1
       se clasifica como REVISAR / CRÍTICO - RIESGO ACTIVO y aparece en la Pestaña 2 (Riesgos_Activos).
    3. Usuario con Active == 1 sin login post-baja se clasifica como REVISAR / ALTO - CUENTA HUÉRFANA en Pestaña 2.
    4. Usuario con reingreso vigente (segundo registro sin fecha de baja) se clasifica como OK / POSIBLE REINGRESO en Pestaña 4.
    """
    df_uni = pd.DataFrame({
        "CLA_TRAB": [1001, 1002, 1003, 1004, 1005, 1005],
        "CLA_LOGIN": ["URES", "UACT", "UHUER", "UCONF", "UREING1", "UREING2"],
        "NOMBRE": [
            "USUARIO RESUELTO",
            "USUARIO ACTIVO",
            "USUARIO HUERFANA",
            "USUARIO CONFORME",
            "USUARIO REINGRESO",
            "USUARIO REINGRESO",
        ],
        "FECHA_BAJA": [
            "01/09/2026",  # USER RESUELTO: baja 01/09/2026, login 15/09/2026, Active 0 -> INCIDENTE RESUELTO
            "01/09/2026",  # USER ACTIVO: baja 01/09/2026, login 15/09/2026, Active 1 -> CRÍTICO - RIESGO ACTIVO
            "01/09/2026",  # USER HUERFANA: baja 01/09/2026, login 20/08/2026, Active 1 -> ALTO - CUENTA HUÉRFANA
            "01/09/2026",  # USER CONFORME: baja 01/09/2026, login 20/08/2026, Active 0 -> CONFORME
            "01/01/2025",  # USER REINGRESO (Registro 1): baja antigua pero segundo registro activo vigente
            None,          # USER REINGRESO (Registro 2): activo vigente sin fecha de baja
        ],
        "CORREO": [
            "resuelto@empresa.com",
            "activo@empresa.com",
            "huerfana@empresa.com",
            "conforme@empresa.com",
            "reingreso@empresa.com",
            "reingreso@empresa.com",
        ],
    })

    df_rep = pd.DataFrame({
        "Full Name": [
            "USUARIO RESUELTO",
            "USUARIO ACTIVO",
            "USUARIO HUERFANA",
            "USUARIO CONFORME",
            "USUARIO REINGRESO",
        ],
        "Username": [
            "resuelto@empresa.com",
            "activo@empresa.com",
            "huerfana@empresa.com",
            "conforme@empresa.com",
            "reingreso@empresa.com",
        ],
        "Active": ["0", "1", "1", "0", "1"],
        "Last Login": [
            "15/09/2026 10:00",
            "15/09/2026 10:00",
            "20/08/2026 10:00",
            "20/08/2026 10:00",
            "15/09/2026 10:00",
        ],
    })

    col_map = {
        "name": "Full Name",
        "user": "Username",
        "active": "Active",
        "login": "Last Login",
        "baja_rep": None,
    }

    consolidated = consolidate_logins(df_rep, col_map)
    df_completo, df_riesgos_activos, metrics = execute_audit(
        df_universo=df_uni,
        baja_col="FECHA_BAJA",
        email_col="CORREO",
        nombre_col="NOMBRE",
        consolidated_logins=consolidated,
    )

    out_file = tmp_path / "Auditoria_Validacion_Criterios.xlsx"
    export_to_excel(df_completo, df_riesgos_activos, out_file)

    # 1. Validación precisa de cada registro
    row_resuelto = df_completo[df_completo["NOMBRE"] == "USUARIO RESUELTO"].iloc[0]
    row_activo = df_completo[df_completo["NOMBRE"] == "USUARIO ACTIVO"].iloc[0]
    row_huerfana = df_completo[df_completo["NOMBRE"] == "USUARIO HUERFANA"].iloc[0]
    row_conforme = df_completo[df_completo["NOMBRE"] == "USUARIO CONFORME"].iloc[0]
    rows_reingreso = df_completo[df_completo["NOMBRE"] == "USUARIO REINGRESO"]

    # Criterio A: login 15/09/2026, baja 01/09/2026 y Active == 0
    assert row_resuelto["ESTATUS_AUDITORIA"] == "INCIDENTE RESUELTO"
    assert row_resuelto["CATEGORIA_RIESGO"] == "MEDIO - INCIDENTE PASADO"
    assert row_resuelto["DIAS_POST_BAJA"] == 14

    # Criterio B: login 15/09/2026, baja 01/09/2026 y Active == 1
    assert row_activo["ESTATUS_AUDITORIA"] == "REVISAR"
    assert row_activo["CATEGORIA_RIESGO"] == "CRÍTICO - RIESGO ACTIVO"
    assert row_activo["DIAS_POST_BAJA"] == 14

    # Criterio C: Active == 1 sin login post-baja
    assert row_huerfana["ESTATUS_AUDITORIA"] == "REVISAR"
    assert row_huerfana["CATEGORIA_RIESGO"] == "ALTO - CUENTA HUÉRFANA"

    # Criterio D: Conforme
    assert row_conforme["ESTATUS_AUDITORIA"] == "OK"
    assert row_conforme["CATEGORIA_RIESGO"] == "CONFORME"

    # Criterio E: Posible Reingreso
    for _, r_reing in rows_reingreso.iterrows():
        assert r_reing["ESTATUS_AUDITORIA"] == "OK"
        assert r_reing["CATEGORIA_RIESGO"] == "POSIBLE REINGRESO"

    # 2. Validación de archivo Excel y 5 pestañas
    assert out_file.exists()
    wb = openpyxl.load_workbook(out_file, data_only=True)
    expected_sheets = [
        "Glosario_y_Criterios",
        "Auditoria_Completa",
        "Riesgos_Activos",
        "Incidentes_Pasados",
        "Reingresos",
    ]
    assert wb.sheetnames == expected_sheets
    assert wb.sheetnames[0] == "Glosario_y_Criterios"

    # Pestaña 2: Riesgos_Activos debe contener USUARIO ACTIVO y USUARIO HUERFANA, pero NO USUARIO RESUELTO
    ws_activos = wb["Riesgos_Activos"]
    nombres_activos = [ws_activos.cell(row=i, column=3).value for i in range(2, ws_activos.max_row + 1)]
    assert "USUARIO ACTIVO" in nombres_activos
    assert "USUARIO HUERFANA" in nombres_activos
    assert "USUARIO RESUELTO" not in nombres_activos

    # Pestaña 3: Incidentes_Pasados debe contener USUARIO RESUELTO, pero NO USUARIO ACTIVO
    ws_pasados = wb["Incidentes_Pasados"]
    nombres_pasados = [ws_pasados.cell(row=i, column=3).value for i in range(2, ws_pasados.max_row + 1)]
    assert "USUARIO RESUELTO" in nombres_pasados
    assert "USUARIO ACTIVO" not in nombres_pasados

    # Pestaña 4: Reingresos debe contener USUARIO REINGRESO
    ws_reing = wb["Reingresos"]
    nombres_reing = [ws_reing.cell(row=i, column=3).value for i in range(2, ws_reing.max_row + 1)]
    assert "USUARIO REINGRESO" in nombres_reing

    wb.close()


def test_caso_4_real_workspace_files(tmp_path):
    """
    Valida la ejecución con los archivos reales del proyecto y verifica
    que las 5 pestañas ('Glosario_y_Criterios', 'Auditoria_Completa', 'Riesgos_Activos',
    'Incidentes_Pasados', 'Reingresos') se generen correctamente.
    """
    uni_path = ROOT_DIR / "Universo_Usuarios_PROD_2.xlsx"
    rep_path = ROOT_DIR / "report1790717536569.xlsx"
    if not rep_path.exists():
        rep_path = ROOT_DIR / "report1790717536569.xls"
    if not rep_path.exists():
        parent_xlsx = ROOT_DIR.parent / "report1790717536569.xlsx"
        if parent_xlsx.exists():
            rep_path = parent_xlsx
        else:
            parent_xls = ROOT_DIR.parent / "report1790717536569.xls"
            if parent_xls.exists():
                rep_path = parent_xls

    if not uni_path.exists() or not rep_path.exists():
        pytest.skip("Archivos reales no disponibles en el entorno.")

    out_file = tmp_path / "Auditoria_Accesos_Test_Suite.xlsx"
    saved_file, metrics = run_audit_pipeline(
        universo_path=uni_path,
        reporte_path=rep_path,
        output_path=out_file,
        hoja_universo="BajasU",
    )

    try:
        assert saved_file.exists()
        assert metrics.total_evaluados == 4955
        # Discrepancias de identidad detectadas por correos compartidos
        assert metrics.total_discrepancias_identidad == 18
        # Los casos de riesgo activo que requieren acción inmediata son 4 (2 crítico + 2 alto huérfana)
        assert metrics.total_revisar == 4
        assert metrics.total_critico_activo == 2
        assert metrics.total_alto_huerfana == 2
        # Los incidentes resueltos / pasados (cuentas ya apagadas) son 15 tras deduplicación de bajas históricas
        assert metrics.total_medio_incidente_pasado == 15
        assert metrics.total_ok == 4936

        wb = openpyxl.load_workbook(saved_file, data_only=True)
        assert wb.sheetnames[0] == "Glosario_y_Criterios"
        assert "Glosario_y_Criterios" in wb.sheetnames
        assert "Auditoria_Completa" in wb.sheetnames
        assert "Riesgos_Activos" in wb.sheetnames
        assert "Incidentes_Pasados" in wb.sheetnames
        assert "Reingresos" in wb.sheetnames

        ws_activos = wb["Riesgos_Activos"]
        # 1 fila encabezado + 4 registros = 5 filas
        assert ws_activos.max_row == 5

        ws_pasados = wb["Incidentes_Pasados"]
        # 1 fila encabezado + 15 registros = 16 filas
        assert ws_pasados.max_row == 16
        wb.close()
    finally:
        if out_file.exists():
            try:
                out_file.unlink()
            except Exception:
                pass


def test_caso_5_glosario_y_criterios_completo():
    """
    Valida exhaustivamente la pestaña 'Glosario_y_Criterios':
    1. Que se inserte en la posición 0 del libro Excel.
    2. Que contenga las 3 secciones obligatorias (Estructura, Diccionario y Reglas Técnicas).
    3. Que contenga la matriz completa de riesgos (CRÍTICO, ALTO, MEDIO, POSIBLE REINGRESO, CONFORME).
    4. Que los estilos visuales, colores de celda y anchos de columna estén configurados.
    5. Que las líneas de cuadrícula estén activas y el texto esté ajustado (wrap_text).
    """
    wb = openpyxl.Workbook()
    ws = create_glosario_sheet(wb)

    # 1. Posición y nombre
    assert wb.sheetnames[0] == "Glosario_y_Criterios"
    assert ws.title == "Glosario_y_Criterios"
    assert ws.views.sheetView[0].showGridLines is True

    # 2. Encabezados de sección
    assert "DICCIONARIO DE AUDITORÍA" in str(ws["A1"].value).upper()
    assert "1. ESTRUCTURA DEL LIBRO" in str(ws["A4"].value).upper()
    assert "2. DICCIONARIO DE ESTATUS" in str(ws["A12"].value).upper()
    assert "3. REGLAS TÉCNICAS" in str(ws["A20"].value).upper()

    # 3. Contenido de Sección 1: Estructura del Libro (Pestañas)
    pestañas_esperadas = [
        "Glosario_y_Criterios",
        "Riesgos_Activos",
        "Incidentes_Pasados",
        "Reingresos",
        "Auditoria_Completa",
    ]
    pestañas_encontradas = [ws.cell(row=r, column=1).value for r in range(6, 11)]
    assert pestañas_encontradas == pestañas_esperadas

    # 4. Contenido de Sección 2: Matriz de Riesgo y Acciones
    riesgos_esperados = [
        "CRÍTICO - RIESGO ACTIVO",
        "ALTO - CUENTA HUÉRFANA",
        "MEDIO - INCIDENTE PASADO",
        "POSIBLE REINGRESO",
        "CONFORME",
    ]
    riesgos_encontrados = [ws.cell(row=r, column=1).value for r in range(14, 19)]
    assert riesgos_encontrados == riesgos_esperados

    estatus_esperados = ["REVISAR", "REVISAR", "INCIDENTE RESUELTO", "OK", "OK"]
    estatus_encontrados = [ws.cell(row=r, column=2).value for r in range(14, 19)]
    assert estatus_encontrados == estatus_esperados

    # Validar condiciones lógicas
    assert "Active == 1" in str(ws.cell(row=14, column=3).value)
    assert "Last Login > Fecha Baja" in str(ws.cell(row=14, column=3).value)
    assert "Login <= Baja" in str(ws.cell(row=15, column=3).value)

    # Validar colores de celdas semánticos
    fill_critico = ws.cell(row=14, column=1).fill.start_color.rgb
    fill_alto = ws.cell(row=15, column=1).fill.start_color.rgb
    fill_medio = ws.cell(row=16, column=1).fill.start_color.rgb
    fill_reingreso = ws.cell(row=17, column=1).fill.start_color.rgb
    fill_conforme = ws.cell(row=18, column=1).fill.start_color.rgb

    assert fill_critico == "FFFEE2E2"
    assert fill_alto == "FFFEF3C7"
    assert fill_medio == "FFF1F5F9"
    assert fill_reingreso == "FFE0F2FE"
    assert fill_conforme == "FFDCFCE7"

    # 5. Contenido de Sección 3: Reglas Técnicas
    reglas_esperadas = ["Corte Calendario", "Priorización de Bajas", "Validación Compuesta"]
    reglas_encontradas = [ws.cell(row=r, column=1).value for r in range(22, 25)]
    assert reglas_encontradas == reglas_esperadas

    # 6. Anchos de columna y wrap_text
    for col_letter in ["A", "B", "C", "D", "E"]:
        assert ws.column_dimensions[col_letter].width >= 20

    assert ws.cell(row=14, column=4).alignment.wrap_text is True
    assert ws.cell(row=14, column=5).alignment.wrap_text is True


def test_fechas_de_baja_multiples():
    """
    Test Fechas de Baja Múltiples:
    - Universo: Usuario X con bajas 15/05/2023 y 10/01/2026.
    - Reporte: Login el 20/06/2024.
    - Esperado: El script debe ignorar la baja de 2023, tomar la de 10/01/2026
      y clasificarlo como OK (el login fue previo a su última baja).
    """
    df_uni = pd.DataFrame({
        "CLA_TRAB": [901, 901],
        "NOMBRE": ["USUARIO MULTI BAJA", "USUARIO MULTI BAJA"],
        "FECHA_BAJA": ["15/05/2023", "10/01/2026"],
        "CORREO": ["multibaja@empresa.com", "multibaja@empresa.com"],
    })

    df_rep = pd.DataFrame({
        "Full Name": ["USUARIO MULTI BAJA"],
        "Username": ["multibaja@empresa.com"],
        "Active": ["0"],
        "Last Login": ["20/06/2024 10:00"],
    })

    col_map = {
        "name": "Full Name",
        "user": "Username",
        "active": "Active",
        "login": "Last Login",
        "baja_rep": None,
    }

    consolidated = consolidate_logins(df_rep, col_map)
    df_completo, df_hallazgos, metrics = execute_audit(
        df_universo=df_uni,
        baja_col="FECHA_BAJA",
        email_col="CORREO",
        nombre_col="NOMBRE",
        consolidated_logins=consolidated,
    )

    assert len(df_completo) == 1
    user = df_completo.iloc[0]
    assert user["FECHA_BAJA"] == "10/01/2026"
    assert user["ULTIMO_LOGIN_DETECTADO"] == "20/06/2024"
    assert user["ESTATUS_AUDITORIA"] == "OK"
    assert user["CATEGORIA_RIESGO"] == "CONFORME"
    assert pd.isna(user["DIAS_POST_BAJA"])
    assert len(df_hallazgos) == 0


def test_login_mas_cercano():
    """
    Test Login Más Cercano:
    - Usuario dado de baja el 01/09/2026.
    - Reporte tiene dos logins: 03/09/2026 (+2 días) y 30/09/2026 (+29 días).
    - Esperado: Debe seleccionar 03/09/2026 por proximidad temporal inmediata al cese.
    """
    df_uni = pd.DataFrame({
        "CLA_TRAB": [902],
        "NOMBRE": ["USUARIO LOGIN CERCANO"],
        "FECHA_BAJA": ["01/09/2026"],
        "CORREO": ["cercano@empresa.com"],
    })

    df_rep = pd.DataFrame({
        "Full Name": ["USUARIO LOGIN CERCANO", "USUARIO LOGIN CERCANO"],
        "Username": ["cercano@empresa.com", "cercano@empresa.com"],
        "Active": ["0", "0"],
        "Last Login": ["30/09/2026 10:00", "03/09/2026 08:30"],
    })

    col_map = {
        "name": "Full Name",
        "user": "Username",
        "active": "Active",
        "login": "Last Login",
        "baja_rep": None,
    }

    consolidated = consolidate_logins(df_rep, col_map)
    df_completo, df_hallazgos, metrics = execute_audit(
        df_universo=df_uni,
        baja_col="FECHA_BAJA",
        email_col="CORREO",
        nombre_col="NOMBRE",
        consolidated_logins=consolidated,
    )

    assert len(df_completo) == 1
    user = df_completo.iloc[0]
    # Debe seleccionar 03/09/2026 (distancia 2 días vs 29 días)
    assert user["ULTIMO_LOGIN_DETECTADO"] == "03/09/2026"
    assert user["DIAS_POST_BAJA"] == 2
    assert user["ESTATUS_AUDITORIA"] == "INCIDENTE RESUELTO"
    assert user["CATEGORIA_RIESGO"] == "MEDIO - INCIDENTE PASADO"


def test_login_empate_distancia_prioriza_posterior():
    """
    Test de Desempate de Distancia:
    - Usuario dado de baja el 01/09/2026.
    - Reporte tiene dos logins equidistantes:
      * 30/08/2026 (-2 días, previo a la baja)
      * 03/09/2026 (+2 días, posterior a la baja)
    - Esperado: Ante empate exacto de distancia, debe priorizar el login posterior
      (03/09/2026) para no omitir un posible acceso no autorizado.
    """
    df_uni = pd.DataFrame({
        "CLA_TRAB": [903],
        "NOMBRE": ["USUARIO EMPATE"],
        "FECHA_BAJA": ["01/09/2026"],
        "CORREO": ["empate@empresa.com"],
    })

    df_rep = pd.DataFrame({
        "Full Name": ["USUARIO EMPATE", "USUARIO EMPATE"],
        "Username": ["empate@empresa.com", "empate@empresa.com"],
        "Active": ["0", "0"],
        "Last Login": ["30/08/2026 10:00", "03/09/2026 10:00"],
    })

    col_map = {
        "name": "Full Name",
        "user": "Username",
        "active": "Active",
        "login": "Last Login",
        "baja_rep": None,
    }

    consolidated = consolidate_logins(df_rep, col_map)
    df_completo, df_hallazgos, metrics = execute_audit(
        df_universo=df_uni,
        baja_col="FECHA_BAJA",
        email_col="CORREO",
        nombre_col="NOMBRE",
        consolidated_logins=consolidated,
    )

    user = df_completo.iloc[0]
    assert user["ULTIMO_LOGIN_DETECTADO"] == "03/09/2026"
    assert user["DIAS_POST_BAJA"] == 2
    assert user["ESTATUS_AUDITORIA"] == "INCIDENTE RESUELTO"
