"""
test_auditor.py
=============================================================================
Suite de pruebas unitarias y de integración para la refactorización de
auditor_bajas.py (Leaver Access Review Automation).
=============================================================================
"""

import sys
from pathlib import Path
import tempfile
import numpy as np
import openpyxl
import pandas as pd
import pytest

# Permitir importación del módulo principal al ejecutar directamente desde tests/
ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from auditor_bajas import (
    clean_identifier_series,
    consolidate_logins,
    create_glosario_sheet,
    execute_audit,
    find_column_by_substring,
    parse_active_series,
    read_reporte_logins,
    read_universo,
    robust_parse_dates,
    run_audit_pipeline,
)


def test_robust_parse_dates():
    """Verifica que fechas en formatos DD/MM/YYYY, timestamps y ceros se parseen adecuadamente."""
    series = pd.Series([
        "03/09/2026 00:00",
        "29/09/2026 03:31 PM",
        "0",
        0,
        "",
        None,
        "nan",
        "2026-09-03",
    ])
    parsed = robust_parse_dates(series)
    assert parsed.iloc[0] == pd.Timestamp("2026-09-03")
    assert parsed.iloc[1] == pd.Timestamp("2026-09-29")
    assert pd.isna(parsed.iloc[2])
    assert pd.isna(parsed.iloc[3])
    assert pd.isna(parsed.iloc[4])
    assert pd.isna(parsed.iloc[5])
    assert pd.isna(parsed.iloc[6])
    assert parsed.iloc[7] == pd.Timestamp("2026-09-03")


def test_parse_active_series():
    """Verifica el parseo robusto del indicador de cuenta activa."""
    series = pd.Series([
        "1",
        " 1 ",
        "0",
        " 0 ",
        1,
        0,
        "true",
        "false",
        "si",
        "no",
        "activo",
        None,
        np.nan,
    ])
    active = parse_active_series(series)
    expected = [1, 1, 0, 0, 1, 0, 1, 0, 1, 0, 1, 0, 0]
    assert active.tolist() == expected


def test_ghost_columns_and_truncated_headers():
    """
    Simula exactamente el problema de columnas vacías desfasadas
    y encabezados cortados ('Fecha de Ba...') en un archivo Excel.
    """
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = Path(tmp.name)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "in"

    # Encabezados con columna vacía intercalada (celda combinada) y 'Fecha de Ba...'
    headers = [
        "Full Name",
        None,  # Columna fantasma vacía
        "Profile",
        "Username",
        "",  # Otra columna fantasma
        "Active",
        "Last Login",
        "Fecha de Ba...",  # Truncado
        "Created Date",
    ]
    ws.append(headers)

    # Filas con datos
    ws.append([
        "JUAN PEREZ", None, "Ventas", "jperez@empresa.com", None, "1", "29/09/2026 10:00", "01/09/2026", "01/01/2025"
    ])
    ws.append([
        "MARIA LOPEZ", None, "Admin", "mlopez@empresa.com", None, "0", "15/08/2026 08:30", "20/08/2026", "10/02/2024"
    ])
    wb.save(tmp_path)

    try:
        df_rep, col_map = read_reporte_logins(tmp_path)
        # Las columnas vacías deben haberse eliminado
        assert None not in df_rep.columns
        assert "" not in df_rep.columns

        # Mapeo correcto de columnas
        assert col_map["user"] == "Username"
        assert col_map["active"] == "Active"
        assert col_map["login"] == "Last Login"
        assert col_map["baja_rep"] == "Fecha de Ba..."
        assert col_map["created"] == "Created Date"
    finally:
        if tmp_path.exists():
            tmp_path.unlink()


def test_multiple_logins_consolidation():
    """Verifica que un usuario repetido conserve el login más reciente y Active=1 si aplica."""
    data = {
        "Username": ["user1@empresa.com", "user1@empresa.com", "user2@empresa.com"],
        "Active": ["0", "1", "0"],
        "Last Login": ["01/08/2026", "25/08/2026", "10/05/2026"],
        "Fecha de Ba...": ["01/07/2026", "01/07/2026", "01/05/2026"],
    }
    df = pd.DataFrame(data)
    col_map = {
        "user": "Username",
        "active": "Active",
        "login": "Last Login",
        "baja_rep": "Fecha de Ba...",
    }

    consolidated = consolidate_logins(df, col_map)
    u1 = consolidated[consolidated["_clean_username"] == "user1@empresa.com"].iloc[0]
    assert u1["_parsed_login_date"] == pd.Timestamp("2026-08-25")
    assert u1["_parsed_active"] == 1

    u2 = consolidated[consolidated["_clean_username"] == "user2@empresa.com"].iloc[0]
    assert u2["_parsed_login_date"] == pd.Timestamp("2026-05-10")
    assert u2["_parsed_active"] == 0


def test_double_risk_audit_matrix():
    """
    Verifica la matriz completa de evaluación de riesgo:
    1. Ambos (Acceso post-baja Y Active == 1)
    2. Acceso Post-Baja (Acceso post-baja Y Active == 0)
    3. Cuenta Activa Huérfana (Active == 1 pero sin login post-baja)
    4. Ninguno (Login <= Baja y Active == 0)
    5. Fallback combine_first de Fecha de Baja desde Reporte
    """
    df_uni = pd.DataFrame({
        "CLA_TRAB": [101, 102, 103, 104, 105, 106],
        "CLA_LOGIN": ["u1", "u2", "u3", "u4", "u5", "u6"],
        "NOMBRE": ["USER AMBOS", "USER ACCESO", "USER ACTIVA", "USER OK", "USER FALLBACK", "USER ACTIVO_EMP"],
        "FECHA_BAJA": [
            "10/08/2026",       # u1: Login 20/08/2026 + Active 1 -> Ambos
            "10/08/2026",       # u2: Login 20/08/2026 + Active 0 -> Acceso Post-Baja
            "10/08/2026",       # u3: Login 05/08/2026 + Active 1 -> Cuenta Activa Huérfana
            "10/08/2026",       # u4: Login 05/08/2026 + Active 0 -> Ninguno / OK
            None,               # u5: Sin fecha en universo, pero viene 05/08/2026 en reporte -> Fallback -> Acceso
            None,               # u6: Sin fecha en universo y sin fecha en reporte -> Empleado activo -> OK
        ],
        "COL_E": [None] * 6,
        "COL_F": [None] * 6,
        "COL_G": [None] * 6,
        "COL_H": [None] * 6,
        "COL_I": [None] * 6,
        "CORREO": [
            "u1@empresa.com",
            "u2@empresa.com",
            "u3@empresa.com",
            "u4@empresa.com",
            "u5@empresa.com",
            "u6@empresa.com",
        ],
    })

    df_rep = pd.DataFrame({
        "Username": [
            "u1@empresa.com",
            "u2@empresa.com",
            "u3@empresa.com",
            "u4@empresa.com",
            "u5@empresa.com",
            "u6@empresa.com",
        ],
        "Active": ["1", "0", "1", "0", "0", "1"],
        "Last Login": [
            "20/08/2026",
            "20/08/2026",
            "05/08/2026",
            "05/08/2026",
            "15/08/2026",
            "29/09/2026",
        ],
        "Fecha de Ba...": [
            "10/08/2026",
            "10/08/2026",
            "10/08/2026",
            "10/08/2026",
            "05/08/2026",  # Rescate de fecha de baja para u5
            None,
        ],
    })

    col_map = {
        "user": "Username",
        "active": "Active",
        "login": "Last Login",
        "baja_rep": "Fecha de Ba...",
    }

    consolidated = consolidate_logins(df_rep, col_map)
    df_completo, df_hallazgos, metrics = execute_audit(
        df_universo=df_uni,
        baja_col="FECHA_BAJA",
        email_col="CORREO",
        consolidated_logins=consolidated,
    )

    # Validaciones en df_completo
    riesgo_map = dict(zip(df_completo["NOMBRE"], df_completo["CATEGORIA_RIESGO"]))
    est_map = dict(zip(df_completo["NOMBRE"], df_completo["ESTATUS_AUDITORIA"]))
    dias_map = dict(zip(df_completo["NOMBRE"], df_completo["DIAS_POST_BAJA"]))

    # USER AMBOS: Acceso post-baja + Active == 1 -> CRÍTICO - RIESGO ACTIVO, REVISAR
    assert riesgo_map["USER AMBOS"] == "CRÍTICO - RIESGO ACTIVO"
    assert est_map["USER AMBOS"] == "REVISAR"
    assert dias_map["USER AMBOS"] == 10

    # USER ACCESO: Acceso post-baja + Active == 0 -> MEDIO - INCIDENTE PASADO, INCIDENTE RESUELTO
    assert riesgo_map["USER ACCESO"] == "MEDIO - INCIDENTE PASADO"
    assert est_map["USER ACCESO"] == "INCIDENTE RESUELTO"
    assert dias_map["USER ACCESO"] == 10

    # USER ACTIVA: Sin acceso post-baja + Active == 1 -> ALTO - CUENTA HUÉRFANA, REVISAR
    assert riesgo_map["USER ACTIVA"] == "ALTO - CUENTA HUÉRFANA"
    assert est_map["USER ACTIVA"] == "REVISAR"
    assert pd.isna(dias_map["USER ACTIVA"])

    # USER OK: Sin acceso post-baja + Active == 0 -> CONFORME, OK
    assert riesgo_map["USER OK"] == "CONFORME"
    assert est_map["USER OK"] == "OK"

    # u5 rescató fecha de baja 05/08/2026 y login 15/08/2026 (Active 0) -> MEDIO - INCIDENTE PASADO
    assert riesgo_map["USER FALLBACK"] == "MEDIO - INCIDENTE PASADO"
    assert est_map["USER FALLBACK"] == "INCIDENTE RESUELTO"
    assert dias_map["USER FALLBACK"] == 10

    # u6 no tiene baja en ninguna fuente -> Empleado activo -> CONFORME, OK
    assert riesgo_map["USER ACTIVO_EMP"] == "CONFORME"
    assert est_map["USER ACTIVO_EMP"] == "OK"

    # Validar métricas
    assert metrics.total_evaluados == 6
    assert metrics.total_revisar == 2
    assert metrics.total_ok == 2
    assert metrics.total_critico_activo == 1
    assert metrics.total_alto_huerfana == 1
    assert metrics.total_medio_incidente_pasado == 2
    assert metrics.total_conforme == 2

    # Validar ordenamiento en df_hallazgos (Riesgos Activos): primero crítico, luego cuenta huérfana
    assert len(df_hallazgos) == 2
    hallazgos_tipos = df_hallazgos["CATEGORIA_RIESGO"].tolist()
    assert hallazgos_tipos[0] == "CRÍTICO - RIESGO ACTIVO"
    assert hallazgos_tipos[-1] == "ALTO - CUENTA HUÉRFANA"


def test_real_workspace_files(tmp_path):
    """Ejecuta el pipeline completo contra los archivos reales en el workspace."""
    uni_path = ROOT_DIR / "Universo_Usuarios_PROD_2.xlsx"
    rep_path = ROOT_DIR / "report1790717536569.xlsx"
    if not rep_path.exists():
        rep_path = ROOT_DIR / "report1790717536569.xls"
    if not rep_path.exists():
        parent_xlsx = ROOT_DIR.parent / "report1790717536569.xlsx"
        if parent_xlsx.exists():
            rep_path = parent_xlsx
        else:
            parent_xls = ROOT_DIR.parent / "report1790717536569.xls"
            if parent_xls.exists():
                rep_path = parent_xls

    if not uni_path.exists() or not rep_path.exists():
        pytest.skip("Archivos reales no disponibles en el entorno.")

    out_file = tmp_path / "Auditoria_Accesos_Resultado_Test.xlsx"
    saved_file, metrics = run_audit_pipeline(
        universo_path=uni_path,
        reporte_path=rep_path,
        output_path=out_file,
        hoja_universo="BajasU",
    )

    try:
        assert saved_file.exists()
        assert metrics.total_evaluados == 4955
        assert metrics.total_revisar == 4
        assert metrics.total_critico_activo == 2
        assert metrics.total_alto_huerfana == 2
        assert metrics.total_medio_incidente_pasado == 15
        assert metrics.total_ok == 4936
        assert metrics.total_discrepancias_identidad == 18

        # Validar estructura y hojas del archivo generado
        wb = openpyxl.load_workbook(saved_file, data_only=True)
        assert wb.sheetnames[0] == "Glosario_y_Criterios"
        assert "Glosario_y_Criterios" in wb.sheetnames
        assert "Auditoria_Completa" in wb.sheetnames
        assert "Riesgos_Activos" in wb.sheetnames
        assert "Incidentes_Pasados" in wb.sheetnames
        assert "Reingresos" in wb.sheetnames

        ws_activos = wb["Riesgos_Activos"]
        # 1 fila encabezado + 4 registros = 5 filas
        assert ws_activos.max_row == 5

        ws_pasados = wb["Incidentes_Pasados"]
        # 1 fila encabezado + 15 registros = 16 filas
        assert ws_pasados.max_row == 16
        wb.close()
    finally:
        if out_file.exists():
            try:
                out_file.unlink()
            except Exception:
                pass


def test_glosario_y_criterios_unitario():
    """Valida la generación aislada de la hoja de Glosario y Criterios."""
    wb = openpyxl.Workbook()
    ws = create_glosario_sheet(wb)

    assert wb.sheetnames[0] == "Glosario_y_Criterios"
    assert ws.title == "Glosario_y_Criterios"
    assert ws.views.sheetView[0].showGridLines is True

    # Validar secciones presentes
    assert "DICCIONARIO DE AUDITORÍA" in str(ws["A1"].value).upper()
    assert "1. ESTRUCTURA DEL LIBRO" in str(ws["A4"].value).upper()
    assert "2. DICCIONARIO DE ESTATUS" in str(ws["A12"].value).upper()
    assert "3. REGLAS TÉCNICAS" in str(ws["A20"].value).upper()

    # Validar categorías críticas y altas
    assert ws["A14"].value == "CRÍTICO - RIESGO ACTIVO"
    assert ws["B14"].value == "REVISAR"
    assert ws["A14"].fill.start_color.rgb == "FFFEE2E2"

    assert ws["A15"].value == "ALTO - CUENTA HUÉRFANA"
    assert ws["B15"].value == "REVISAR"
    assert ws["A15"].fill.start_color.rgb == "FFFEF3C7"



def test_config_y_carpeta_resultados(tmp_path, monkeypatch):
    import auditor_bajas as ab

    monkeypatch.setenv("AUDITORIA_CONFIG", str(tmp_path / "cfg" / "config.json"))
    assert ab.load_config() == {}
    ab.save_config({"output_dir": "X"})
    assert ab.load_config() == {"output_dir": "X"}

    # Primera vez: ruta dentro de scripts rechazada, luego una valida que se crea tras confirmar
    destino = tmp_path / "resultados"
    respuestas = iter([str(ab.SCRIPT_DIR / "res"), str(destino), "s"])
    monkeypatch.setattr(ab, "_ask", lambda prompt: next(respuestas))
    assert ab.resolve_output_dir({}) == destino
    assert destino.is_dir()

    # Siguientes veces: confirmar la ultima ruta o cambiarla
    monkeypatch.setattr(ab, "_ask", lambda prompt: "s")
    assert ab.resolve_output_dir({"output_dir": str(destino)}) == destino
    nuevo = tmp_path / "otra"
    respuestas = iter(["n", str(nuevo), "s"])
    monkeypatch.setattr(ab, "_ask", lambda prompt: next(respuestas))
    assert ab.resolve_output_dir({"output_dir": str(destino)}) == nuevo
